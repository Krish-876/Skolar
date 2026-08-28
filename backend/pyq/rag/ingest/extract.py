import logging
import os
from pathlib import Path

import pdfplumber

logger = logging.getLogger(__name__)

# ── Step 1 — PDF Parsing ──────────────────────────────────────────────────────
#
# Hybrid per-page OCR: previously OCR only ran when the WHOLE document
# extracted to under 200 chars. That misses the common case of a page that
# has real text plus one question pasted in as a screenshot/diagram — the
# page-level text looks "fine" so OCR never triggered and that question
# silently vanished. Now any page that embeds a raster image AND whose
# extracted text is thin gets OCR'd individually and the result appended,
# tagged so the LLM extraction step can still pick it up. Downstream
# _deduplicate_questions already collapses near-duplicates, so being
# generous here is safe.

def extract_raw_text(pdf_bytes: bytes) -> str:
    import io
    import unicodedata

    def _clean(text: str) -> str:
        text = unicodedata.normalize("NFKC", text)
        return "\n".join(
            line for line in text.splitlines()
            if line.strip()
        )

    full_text = ""
    pages_needing_ocr: list[int] = []

    with pdfplumber.open(io.BytesIO(pdf_bytes)) as pdf:
        for page in pdf.pages:
            # 1. Normal text layer
            text = page.extract_text(x_tolerance=3, y_tolerance=3)

            # 2. Table-aware fallback
            if not text or len(text.strip()) < 40:
                tables = page.extract_tables()
                if tables:
                    rows = []
                    for table in tables:
                        for row in table:
                            rows.append("  |  ".join(cell or "" for cell in row))
                    text = "\n".join(rows)

            # 3. Words-based reconstruction
            if not text or len(text.strip()) < 40:
                words = page.extract_words(
                    x_tolerance=3,
                    y_tolerance=3,
                    keep_blank_chars=False,
                    use_text_flow=True,
                )
                if words:
                    text = " ".join(w["text"] for w in words)

            if text:
                full_text += _clean(text) + "\n"

            # 4. Flag pages with embedded images + thin text for page-level OCR.
            #    Threshold of 400 chars is deliberately generous — a page with
            #    a full paragraph of real text AND a small screenshot question
            #    below it can still clear this and get OCR'd too.
            has_images = bool(page.images)
            text_len = len(text.strip()) if text else 0
            if has_images and text_len < 400:
                pages_needing_ocr.append(page.page_number)

    for page_num in pages_needing_ocr:
        ocr_text = _ocr_single_page(pdf_bytes, page_num)
        if ocr_text.strip():
            full_text += f"\n[OCR supplement — page {page_num}, likely image-based content]\n"
            full_text += _clean(ocr_text) + "\n"

    # 5. OCR fallback for image-only PDFs (whole doc still thin after the above)
    if len(full_text.strip()) < 200:
        full_text = _ocr_fallback(pdf_bytes)

    return full_text


_easyocr_reader = None


def _get_easyocr_reader():
    global _easyocr_reader
    if _easyocr_reader is None:
        import easyocr
        logger.info("[ocr] loading EasyOCR reader (downloads models on first run)")
        _easyocr_reader = easyocr.Reader(["en"], gpu=False)
    return _easyocr_reader


def _ocr_pil_image(img) -> str:
    """OCR a PIL Image via EasyOCR. Returns "" on any failure rather than raising,
    so a bad single page/image never takes down the whole extraction run."""
    try:
        import numpy as np
        reader = _get_easyocr_reader()
        lines = reader.readtext(np.array(img), detail=0, paragraph=True)
        return "\n".join(lines)
    except ImportError:
        logger.warning("[ocr] easyocr not installed — run: pip install easyocr")
        return ""
    except Exception as e:
        logger.warning(f"[ocr] failed: {e}")
        return ""


def _render_pdf_page(pdf_bytes: bytes, page_number: int, dpi: int = 300):
    """Render one 1-indexed PDF page to a PIL Image via pypdfium2."""
    import pypdfium2 as pdfium
    pdf = pdfium.PdfDocument(pdf_bytes)
    try:
        page = pdf[page_number - 1]
        try:
            bitmap = page.render(scale=dpi / 72)
            return bitmap.to_pil()
        finally:
            page.close()
    finally:
        pdf.close()


def _ocr_single_page(pdf_bytes: bytes, page_number: int) -> str:
    """OCR one specific page (1-indexed) of a PDF."""
    try:
        img = _render_pdf_page(pdf_bytes, page_number)
        return _ocr_pil_image(img)
    except ImportError:
        logger.warning("[ocr_single_page] pypdfium2 not installed — run: pip install pypdfium2")
        return ""
    except Exception as e:
        logger.warning(f"[ocr_single_page] page {page_number} failed: {e}")
        return ""


def _ocr_fallback(pdf_bytes: bytes) -> str:
    """OCR every page of a PDF — used when the whole document is text-sparse."""
    try:
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(pdf_bytes)
        try:
            pages_text = []
            for i in range(len(pdf)):
                page = pdf[i]
                try:
                    bitmap = page.render(scale=300 / 72)
                    img = bitmap.to_pil()
                finally:
                    page.close()
                text = _ocr_pil_image(img)
                if text.strip():
                    pages_text.append(text.strip())
            return "\n".join(pages_text)
        finally:
            pdf.close()
    except ImportError:
        logger.warning(
            "[ocr_fallback] pypdfium2/easyocr not installed — "
            "image-only PDFs will extract empty. "
            "Run: pip install pypdfium2 easyocr"
        )
        return ""
    except Exception as e:
        logger.warning(f"[ocr_fallback] failed: {e}")
        return ""

# ── Step 1a — DOCX / DOC / image parsing ──────────────────────────────────────
#
# Adds three more input formats on top of PDF:
#   .docx  — python-docx for paragraphs/tables, + OCR of embedded images
#   .doc   — legacy binary format, not readable by python-docx. Converted to
#            .docx via a local LibreOffice install (free, but a real binary
#            install — not a pip package), then routed through extract_docx_text.
#   images — a standalone photo/screenshot of a question paper, OCR'd directly.

def extract_docx_text(docx_bytes: bytes) -> str:
    """Extract text from a .docx: paragraphs, tables, and OCR of embedded images."""
    import io
    import zipfile

    import docx as python_docx

    full_text: list[str] = []

    doc = python_docx.Document(io.BytesIO(docx_bytes))
    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text.strip())

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                full_text.append("  |  ".join(cells))

    # docx is a zip archive — embedded images live under word/media/
    try:
        with zipfile.ZipFile(io.BytesIO(docx_bytes)) as z:
            media_files = [n for n in z.namelist() if n.startswith("word/media/")]
            if media_files:
                import pytesseract
                from PIL import Image

                for name in media_files:
                    try:
                        img_bytes = z.read(name)
                        img = Image.open(io.BytesIO(img_bytes))
                        ocr_text = pytesseract.image_to_string(img, config="--psm 6").strip()
                        if ocr_text:
                            full_text.append(f"[OCR — embedded image {name}]")
                            full_text.append(ocr_text)
                    except Exception as e:
                        logger.warning(f"[extract_docx_text] failed to OCR {name}: {e}")
    except ImportError:
        logger.warning(
            "[extract_docx_text] pytesseract/Pillow not installed — "
            "embedded images in this docx will be skipped."
        )
    except zipfile.BadZipFile as e:
        logger.warning(f"[extract_docx_text] could not read docx as zip: {e}")

    return "\n".join(full_text)

def extract_doc_text(doc_bytes: bytes, filename: str = "input.doc") -> str:
    """
    Legacy .doc → convert to .docx via local LibreOffice, then extract.
    Requires `soffice` (LibreOffice) installed and on PATH.
    Raises a clear error if it's missing rather than silently returning "".
    """
    import shutil
    import subprocess
    import tempfile

    soffice_path = shutil.which("soffice")
    if soffice_path is None:
        raise RuntimeError(
            "LibreOffice ('soffice') not found on PATH. Legacy .doc files need "
            "it for conversion. Install LibreOffice (free, libreoffice.org) and "
            "make sure 'soffice' is callable from cmd.exe, or convert the .doc "
            "to .docx/.pdf yourself before uploading."
        )

    with tempfile.TemporaryDirectory() as tmp:
        src_path = str(Path(tmp) / Path(filename).name)
        with open(src_path, "wb") as f:
            f.write(doc_bytes)

        result = subprocess.run(  # noqa: S603 -- soffice_path resolved via shutil.which; args are not shell-interpreted
            [soffice_path, "--headless", "--convert-to", "docx", "--outdir", tmp, src_path],
            capture_output=True, text=True, timeout=120,
            check=False,
        )
        if result.returncode != 0:
            raise RuntimeError(f"soffice conversion failed: {result.stderr.strip()}")

        converted_path = os.path.splitext(src_path)[0] + ".docx"
        if not os.path.exists(converted_path):
            raise RuntimeError("soffice did not produce the expected .docx output.")

        with open(converted_path, "rb") as f:
            docx_bytes = f.read()

    return extract_docx_text(docx_bytes)


def extract_image_text(image_bytes: bytes) -> str:
    """OCR a standalone image file (photo/screenshot of a question paper)."""
    try:
        import io

        from PIL import Image
        img = Image.open(io.BytesIO(image_bytes))
        return _ocr_pil_image(img)
    except Exception as e:
        logger.warning(f"[extract_image_text] failed: {e}")
        return ""


_TEXT_EXTRACTORS = {
    ".pdf":  lambda b, fn: extract_raw_text(b),
    ".docx": lambda b, fn: extract_docx_text(b),
    ".doc":  lambda b, fn: extract_doc_text(b, fn),
    ".png":  lambda b, fn: extract_image_text(b),
    ".jpg":  lambda b, fn: extract_image_text(b),
    ".jpeg": lambda b, fn: extract_image_text(b),
    ".webp": lambda b, fn: extract_image_text(b),
    ".tif":  lambda b, fn: extract_image_text(b),
    ".tiff": lambda b, fn: extract_image_text(b),
}


def extract_raw_text_any(file_bytes: bytes, filename: str) -> str:
    """Dispatch to the right extractor based on file extension."""
    ext = os.path.splitext(filename)[1].lower()
    extractor = _TEXT_EXTRACTORS.get(ext)
    if extractor is None:
        raise ValueError(
            f"Unsupported file type '{ext}' for {filename}. "
            f"Supported: {', '.join(sorted(_TEXT_EXTRACTORS))}"
        )
    return extractor(file_bytes, filename)
